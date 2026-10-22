"""Format a Pandoc report and QA the installed Word render; no model execution."""
import hashlib,json,sys
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pymupdf

ROOT=Path(__file__).resolve().parents[2]
BASE=ROOT/'reports/PROJECT_REVIEW_AND_NEXT_PLAN'
EVIDENCE=ROOT/'reports/project_review_evidence'
QA=ROOT/'tmp/project_review_qa'
EVIDENCE.mkdir(exist_ok=True);QA.mkdir(exist_ok=True)
def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
receipt=EVIDENCE/'document_checks.json'
if '--record-visual' in sys.argv:
    r=json.loads(receipt.read_text(encoding='utf-8'))
    assert all(sha(BASE.with_suffix('.'+k))==v for k,v in r['sha256'].items())
    r.update(visual_inspection='PASS',pages_inspected=list(range(1,r['pages']+1)))
    receipt.write_text(json.dumps(r,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print('PASS: current report visually checked');sys.exit()

d=Document(BASE.with_suffix('.docx'))
for name in ['Normal','Body Text','First Paragraph']:
    if name not in d.styles:continue
    style=d.styles[name];style.font.size=Pt(10.5)
    style.font.color.rgb=RGBColor(0,0,0)
    style.paragraph_format.line_spacing=1.15
    style.paragraph_format.space_after=Pt(6)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
for style in d.styles:
    if style.name.startswith('Heading') or style.name in ['Title','Subtitle']:
        style.font.color.rgb=RGBColor(0,0,0)
        style.font.underline=False
        for n in list(style.element.iter(qn('w:pBdr'))):n.getparent().remove(n)
d.styles['Title'].font.size=Pt(22)
for p in d.paragraphs:
    if p.style.name=='Heading 1':
        p.paragraph_format.keep_with_next=True
        if p.text.startswith(('4 ','5 ','7 ','9 ')):p.paragraph_format.page_break_before=True
    if p.style.name in ['Title','Subtitle'] or p.style.name.startswith('Heading'):
        for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
        for n in list(p._p.iter(qn('w:pBdr'))):n.getparent().remove(n)
    if p._p.xpath('.//w:drawing'):
        p.paragraph_format.keep_with_next=True
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
for table in d.tables:
    table.autofit=False
    ncols=len(table.columns)
    widths={3:[3.0,6.6,7.0],4:[1.2,5.3,6.7,3.4],6:[4.5,2.42,2.42,2.42,2.42,2.42],2:[4.7,11.9]}.get(ncols)
    if widths:
        for col,w in zip(table.columns,widths):col.width=Cm(w)
    props=table._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge)
        for k,v in {'val':'single','sz':'4','color':'D9D9D9'}.items():e.set(qn('w:'+k),v)
        borders.append(e)
    props.append(borders)
    for i,row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if i==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for j,cell in enumerate(row.cells):
            if widths:cell.width=Cm(widths[j])
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'17365D' if i==0 else ('F2F5F8' if i%2 else 'FFFFFF'))
            cell._tc.get_or_add_tcPr().append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next=(len(table.rows)<=7 and i<len(table.rows)-1)
                p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
                p.paragraph_format.line_spacing=1.05
                for r in p.runs:
                    r.font.size=Pt(9.5);r.font.bold=(i==0)
                    r.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(0,0,0)
for section in d.sections:
    section.top_margin=Cm(1.9);section.bottom_margin=Cm(1.9)
    section.left_margin=Cm(2.0);section.right_margin=Cm(2.0)
    for container in [section.header,section.footer]:
        for p in container.paragraphs:
            for r in p.runs:
                r.text=r.text.replace('M3 概念参数化几何','项目总结与后续安排')
                r.font.color.rgb=RGBColor(0,0,0)
d.core_properties.title='驱制融合热耦合研究项目总结与后续工作安排'
d.core_properties.language='zh-CN';d.save(BASE.with_suffix('.docx'))
with ZipFile(BASE.with_suffix('.docx')) as z:
    x=ET.fromstring(z.read('word/document.xml'))
    math=len(x.findall('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath'))
    imgs=len([n for n in z.namelist() if n.startswith('word/media/')])
assert math>=3 and imgs==2,(math,imgs)

import win32com.client
app=win32com.client.DispatchEx('Word.Application')
app.Visible=False;app.DisplayAlerts=0;app.AutomationSecurity=3;doc=None
try:
    doc=app.Documents.Open(str(BASE.with_suffix('.docx')),ReadOnly=True,AddToRecentFiles=False)
    doc.ExportAsFixedFormat(str(BASE.with_suffix('.pdf')),17)
finally:
    if doc is not None:doc.Close(False)
    app.Quit()
pdf=pymupdf.open(BASE.with_suffix('.pdf'))
for i,page in enumerate(pdf,1):page.get_pixmap(matrix=pymupdf.Matrix(1.1,1.1)).save(QA/f'page_{i:02d}.png')
text='\n'.join(p.get_text() for p in pdf)
for token in ['PARTIAL','零再生','NR-P','0.001957','PLANNED','SSRN']:assert token in text,token
assert 'PENDING' not in BASE.with_suffix('.md').read_text(encoding='utf-8'), 'Draft sentinel remains'
inputs=['docs/STAGE_VERSIONS.md','docs/NO_REGEN_WORK_PLAN.md','docs/REAR_DRIVE_V04_SOURCED_BASELINE.md',
 'reports/M0_report.md','reports/M1_report.md','reports/M2_report.md','reports/M3_helical_revision.md',
 'reports/M3_APPROVED_SCOPE_2026-10-21.md','reports/M4_report.md','results/m4/acceptance_checks.json',
 'results/m4/studies/D_normal12_rd95_dt1/study_settings.json','results/m4/studies/D_normal12_rd95_dt1/source_intervals.csv',
 'results/m3_helical/assembly_helical.png','results/m4/refined_source_scenarios.png']
(EVIDENCE/'source_manifest.json').write_text(json.dumps(dict(evidence_head='d4c06106c96ce07f7ff691a6b7b9ce8b58dd2f2d',
    files={n:sha(ROOT/n) for n in inputs},scope='Existing evidence review and limited literature verification; no thermal model executed.'),ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
r=dict(status='PASS_STRUCTURE_AND_WORD_RENDER',pages=len(pdf),editable_math_objects=math,figures=imgs,
    tables=len(d.tables),visual_inspection='PENDING',renderer='Installed Word COM ExportAsFixedFormat plus PyMuPDF page images',
    skill_renderer_limitation='Bundled render_docx.py unavailable: pdf2image missing; Word native rendering used.',
    sha256={e:sha(BASE.with_suffix('.'+e)) for e in ['md','docx','pdf']})
receipt.write_text(json.dumps(r,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(r)
