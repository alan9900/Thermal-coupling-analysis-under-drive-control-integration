"""Pandoc DOCX structure plus actual Word PDF render; visual pass is manual."""
import sys,json,hashlib
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET
import pymupdf
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'results/m4';BASE='M4_report'

def main():
    path=ROOT/'reports'/f'{BASE}.docx';receipt=OUT/'document_checks.json'
    if '--record-visual' in sys.argv:
        r=json.loads(receipt.read_text(encoding='utf-8-sig'))
        for ext,h in r['sha256'].items():
            assert hashlib.sha256((ROOT/'reports'/f'{BASE}.{ext}').read_bytes()).hexdigest()==h
        r.update(visual_inspection='PASS',pages_inspected=list(range(1,r['pdf_pages']+1)))
        receipt.write_text(json.dumps(r,indent=2)+'\n');return
    d=Document(path)
    for table in d.tables:
        for i,row in enumerate(table.rows):
            props=row._tr.get_or_add_trPr();props.append(OxmlElement('w:cantSplit'))
            if i==0:props.append(OxmlElement('w:tblHeader'))
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after=Pt(3)
                    # Do not force a long table to stay on a single page.
                    paragraph.paragraph_format.keep_with_next=(i<len(table.rows)-1 if len(table.rows)<=10 else i==0)
                    for run in paragraph.runs:run.font.size=Pt(8.5)
    for section in d.sections:
        for container in [section.header,section.footer]:
            for paragraph in container.paragraphs:
                for run in paragraph.runs:
                    if 'M3 概念参数化几何' in run.text:
                        run.text=run.text.replace('M3 概念参数化几何','M4 Case0 瞬态热分析')
    d.core_properties.language='zh-CN';d.save(path)
    with ZipFile(path) as z:
        xml=ET.fromstring(z.read('word/document.xml'))
        math=len(xml.findall('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath'))
        images=len([n for n in z.namelist() if n.startswith('word/media/')])
        assert math>=6 and images>=4,(math,images)
    import win32com.client
    app=win32com.client.DispatchEx('Word.Application')
    app.Visible=False;app.DisplayAlerts=0;app.AutomationSecurity=3;doc=None
    try:
        doc=app.Documents.Open(str(path),ReadOnly=True,AddToRecentFiles=False)
        doc.ExportAsFixedFormat(str(ROOT/'reports'/f'{BASE}.pdf'),17)
    finally:
        if doc is not None:doc.Close(False)
        app.Quit()
    pdf=pymupdf.open(ROOT/'reports'/f'{BASE}.pdf')
    folder=ROOT/'tmp/m4_report_qa';folder.mkdir(exist_ok=True)
    for i,page in enumerate(pdf):
        page.get_pixmap(matrix=pymupdf.Matrix(1.1,1.1)).save(folder/f'page_{i+1:02d}.png')
    text='\n'.join(p.get_text() for p in pdf)
    for token in ['Case0','S3','PARTIAL','273.15']:assert token in text,token
    r=dict(status='PASS_RENDER_STRUCTURE',editable_math_objects=math,figures=images,pdf_pages=len(pdf),
        visual_inspection='PENDING',source='Markdown through local Pandoc skill; Word COM ExportAsFixedFormat',
        sha256={ext:hashlib.sha256((ROOT/'reports'/f'{BASE}.{ext}').read_bytes()).hexdigest() for ext in ['md','docx','pdf']})
    receipt.write_text(json.dumps(r,indent=2)+'\n');print(r)
if __name__=='__main__':main()
