"""Format Pandoc output and render through installed Word for visual QA."""
import hashlib,json,sys
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pymupdf
ROOT=Path(__file__).resolve().parents[2]
BASE=ROOT/'reports/M4_contour_explained'
QA=ROOT/'tmp/m4_contour_book_qa'
RECEIPT=ROOT/'results/m4/contour_captures/document_checks.json'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
if '--record-visual' in sys.argv:
    r=json.loads(RECEIPT.read_text(encoding='utf-8'))
    assert all(sha(BASE.with_suffix('.'+e))==h for e,h in r['sha256'].items())
    r.update(visual_inspection='PASS',pages_inspected=list(range(1,r['pages']+1)))
    RECEIPT.write_text(json.dumps(r,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print('All current pages visually inspected');sys.exit()
d=Document(BASE.with_suffix('.docx'))
normal=d.styles['Normal'];normal.font.size=Pt(11)
normal.paragraph_format.line_spacing=1.15
normal.paragraph_format.space_after=Pt(7)
normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
title=d.styles['Title'];title.font.color.rgb=RGBColor(0,0,0);title.font.size=Pt(21)
title.font.underline=False
for node in list(title.element.iter(qn('w:pBdr'))):node.getparent().remove(node)
for p in d.paragraphs:
    if p.style.name=='Title':
        for node in list(p._p.iter(qn('w:pBdr'))):node.getparent().remove(node)
        for run in p.runs:run.font.color.rgb=RGBColor(0,0,0);run.font.underline=False
    if p.style.name=='Heading 1' and p.text.startswith('图'):
        p.paragraph_format.page_break_before=True
        p.paragraph_format.keep_with_next=True
    if p._p.xpath('.//w:drawing'):
        p.paragraph_format.space_after=Pt(6)
        p.paragraph_format.keep_with_next=True
for section in d.sections:
    section.top_margin=Cm(1.9);section.bottom_margin=Cm(1.9)
    section.left_margin=Cm(2.0);section.right_margin=Cm(2.0)
    for container in [section.header,section.footer]:
        for p in container.paragraphs:
            for run in p.runs:
                if 'M3 概念参数化几何' in run.text:
                    run.text=run.text.replace('M3 概念参数化几何','M4 Case0 温度云图讲解')
d.core_properties.title='M4 Case0 整体与各部件温度云图讲解'
d.core_properties.language='zh-CN';d.save(BASE.with_suffix('.docx'))
import win32com.client
app=win32com.client.DispatchEx('Word.Application');app.Visible=False;app.DisplayAlerts=0;app.AutomationSecurity=3
doc=None
try:
    doc=app.Documents.Open(str(BASE.with_suffix('.docx')),ReadOnly=True,AddToRecentFiles=False)
    doc.ExportAsFixedFormat(str(BASE.with_suffix('.pdf')),17)
finally:
    if doc is not None:doc.Close(False)
    app.Quit()
QA.mkdir(exist_ok=True)
pdf=pymupdf.open(BASE.with_suffix('.pdf'))
with ZipFile(BASE.with_suffix('.docx')) as z:
    image_count=len([n for n in z.namelist() if n.startswith('word/media/')])
assert image_count==29,image_count
fig_pages=[]
for i,page in enumerate(pdf,1):
    text=page.get_text()
    page.get_pixmap(matrix=pymupdf.Matrix(1.1,1.1)).save(QA/f'page_{i:02d}.png')
    if '区域标识' in text:
        assert len(page.get_images())==1,(i,'one native contour per component page')
        fig_pages.append(i)
assert len(fig_pages)==29,fig_pages
r=dict(status='PASS_STRUCTURE_AND_WORD_RENDER',pages=len(pdf),native_contour_images=image_count,
    figure_pages=fig_pages,visual_inspection='PENDING',source='Pandoc Markdown; Word native PDF render',
    sha256={e:sha(BASE.with_suffix('.'+e)) for e in ['md','docx','pdf']})
RECEIPT.write_text(json.dumps(r,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(r)
