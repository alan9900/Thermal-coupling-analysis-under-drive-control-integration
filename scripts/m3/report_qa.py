"""Pandoc DOCX reference, editable-math inspection and Word PDF rendering."""
import hashlib,json,sys,xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZipFile
import pymupdf
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
ROOT=Path(__file__).resolve().parents[2]

def reference():
    d=Document(ROOT/'reports/M2_reference.docx')
    for s in d.sections:
        for p in s.footer.paragraphs:
            for run in p.runs:run.text=run.text.replace('M2 概念瞬态热网络','M3 概念参数化几何')
    d.save(ROOT/'reports/M3_reference.docx')

def render():
    import win32com.client
    path=ROOT/'reports/M3_report.docx';d=Document(path)
    for table in d.tables:
        for i,row in enumerate(table.rows):
            tr=row._tr.get_or_add_trPr();tr.append(OxmlElement('w:cantSplit'))
            if i==0:tr.append(OxmlElement('w:tblHeader'))
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.05
                    for r in p.runs:r.font.size=Pt(9.5);r.font.bold=(i==0)
    d.core_properties.language='zh-CN';d.save(path)
    with ZipFile(path) as z:
        tree=ET.fromstring(z.read('word/document.xml'));math=len(tree.findall('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath'))
        figures=len([n for n in z.namelist() if n.startswith('word/media/')]);assert math>=12 and figures==5,(math,figures)
    word=win32com.client.DispatchEx('Word.Application');word.Visible=False;word.DisplayAlerts=0;word.AutomationSecurity=3;doc=None
    try:
        doc=word.Documents.Open(str(path),ReadOnly=True,AddToRecentFiles=False);doc.ExportAsFixedFormat(str(ROOT/'reports/M3_report.pdf'),17)
    finally:
        if doc is not None:doc.Close(False)
        word.Quit()
    pdf=pymupdf.open(ROOT/'reports/M3_report.pdf');text='\n'.join(p.get_text() for p in pdf)
    assert all(s in text for s in ['M2-v1.0','0.482','729','Case0','NOT EXECUTED'])
    folder=ROOT/'tmp/m3_qa';folder.mkdir(exist_ok=True,parents=True)
    for i,p in enumerate(pdf):p.get_pixmap(matrix=pymupdf.Matrix(1.3,1.3)).save(folder/f'page_{i+1:02d}.png')
    r=dict(status='PASS',word_export='PASS',editable_math_objects=math,pdf_pages=len(pdf),figures=figures,visual_inspection='PENDING',
           document_sha256={n:hashlib.sha256((ROOT/'reports'/n).read_bytes()).hexdigest() for n in ['M3_report.md','M3_report.docx','M3_report.pdf']},
           post_pandoc_changes='Formatting only: repeating table headers, non-splitting rows, font spacing, language.')
    (ROOT/'results/m3/document_checks.json').write_text(json.dumps(r,indent=2)+'\n',encoding='utf-8');print(json.dumps(r,indent=2))

def record_visual():
    path=ROOT/'results/m3/document_checks.json';r=json.loads(path.read_text())
    for n,s in r['document_sha256'].items():assert hashlib.sha256((ROOT/'reports'/n).read_bytes()).hexdigest()==s
    r.update(visual_inspection='PASS',visually_inspected_pages=list(range(1,r['pdf_pages']+1)),review='All rendered pages inspected after actual Word export; equations, five figures, tables, pagination and Chinese text checked.')
    path.write_text(json.dumps(r,indent=2)+'\n',encoding='utf-8');print('Visual review recorded; hashes unchanged.')

if __name__=='__main__':{'reference':reference,'render':render,'record-visual':record_visual}[sys.argv[1]]()
