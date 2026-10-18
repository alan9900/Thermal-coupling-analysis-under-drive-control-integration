"""Create Pandoc reference styles, or render the resulting DOCX using Word."""
import json
import sys
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET

import pymupdf
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
REPORT = ROOT / "reports/M0_report.docx"
TMP = ROOT / "tmp/m0_qa/report"


def reference():
    d=Document()
    sec=d.sections[0]
    sec.page_height=Cm(29.7); sec.page_width=Cm(21)
    sec.top_margin=Cm(2.0); sec.bottom_margin=Cm(1.9)
    sec.left_margin=Cm(2.1); sec.right_margin=Cm(2.1)
    for name in ["Normal","Body Text","First Paragraph","Caption","Table"]:
        if name not in d.styles:
            continue
        style=d.styles[name]
        style.font.name="Calibri"; style.font.size=Pt(10)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),"宋体")
        style.paragraph_format.space_after=Pt(5)
        style.paragraph_format.line_spacing=1.15
    for name in ["Title","Subtitle","Heading 1","Heading 2","Heading 3"]:
        st=d.styles[name]; st.font.name="Calibri"
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),"黑体")
        st.font.color.rgb=RGBColor.from_string("1E3A5F")
        st.paragraph_format.keep_with_next=True
    d.styles['Title'].font.size=Pt(22)
    d.styles['Heading 1'].font.size=Pt(15)
    d.styles['Heading 2'].font.size=Pt(12)
    d.styles['Heading 1'].paragraph_format.space_before=Pt(14)
    for name in ["Source Code","Verbatim Char"]:
        if name in d.styles:
            d.styles[name].font.size=Pt(8)
    p=sec.footer.paragraphs[0]; p.alignment=2
    p.add_run("V0.4 · M0 参数审查   |   ")
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); p._p.append(fld)
    d.save(ROOT / "reports/M0_reference.docx")


def render():
    import win32com.client
    TMP.mkdir(parents=True,exist_ok=True)
    # Formatting only after Pandoc: keep rows intact and make headers legible.
    # Content and OMML math remain those authored in Markdown.
    styled=Document(REPORT)
    for table in styled.tables:
        for ri,row in enumerate(table.rows):
            trpr=row._tr.get_or_add_trPr()
            if trpr.find(qn('w:cantSplit')) is None:
                trpr.append(OxmlElement('w:cantSplit'))
            for cell in row.cells:
                tcpr=cell._tc.get_or_add_tcPr()
                if ri==0:
                    shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'EAF0F7'); tcpr.append(shade)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before=Pt(0)
                    p.paragraph_format.space_after=Pt(3)
                    p.paragraph_format.line_spacing=1.05
                    for run in p.runs:
                        run.font.size=Pt(9.5)
                        if ri==0: run.font.bold=True
    styled.core_properties.language='zh-CN'
    styled.save(REPORT)
    with ZipFile(REPORT) as z:
        tree=ET.fromstring(z.read('word/document.xml'))
        namespaces={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math',
                    'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        math_count=len(tree.findall('.//m:oMath',namespaces))
        tables=len(tree.findall('.//w:tbl',namespaces))
        images=len([n for n in z.namelist() if n.startswith('word/media/')])
        assert math_count>=20 and tables>=8 and images>=1
    word=win32com.client.DispatchEx('Word.Application')
    word.Visible=False; word.DisplayAlerts=0; word.AutomationSecurity=3
    doc=None
    try:
        doc=word.Documents.Open(str(REPORT),ReadOnly=True,AddToRecentFiles=False)
        doc.ExportAsFixedFormat(str(ROOT / 'reports/M0_report.pdf'),17)
    finally:
        if doc is not None: doc.Close(False)
        word.Quit()
    pdf=pymupdf.open(ROOT / 'reports/M0_report.pdf')
    text='\n'.join(p.get_text() for p in pdf)
    assert '752.84' in text and '149.53' in text and '13500' in text
    for i,p in enumerate(pdf):
        p.get_pixmap(matrix=pymupdf.Matrix(1.25,1.25)).save(TMP / f'page_{i+1:02d}.png')
    result=dict(docx_editable_omml_objects=math_count, tables=tables, embedded_images=images,
                post_pandoc_changes='formatting only: table headers, row integrity, spacing, language metadata',
                pdf_pages=len(pdf), word_pdf_export='PASS', visual_inspection='PENDING')
    (ROOT/'results/m0/document_checks.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
    print(json.dumps(result,indent=2))


if __name__=='__main__':
    {'reference':reference,'render':render}[sys.argv[1]]()
