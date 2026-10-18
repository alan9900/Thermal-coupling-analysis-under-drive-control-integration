"""M0 evidence/configuration audit only. Does not execute EPA/REVS code."""
import csv
import hashlib
import json
import platform
import re
import subprocess
from pathlib import Path
from zipfile import ZipFile

import numpy as np
import yaml
from openpyxl import load_workbook
from scipy.interpolate import RegularGridInterpolator
from scipy.io import savemat

ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / "data/raw/m0"
OUT = ROOT / "results/m0"


def write_json(path, value):
    path.write_text(json.dumps(value, indent=2, ensure_ascii=False, allow_nan=False), encoding="utf-8")


def array_assignment(text, name):
    match = re.search(re.escape(name) + r"\s*=\s*\[([\s\S]*?)\]\s*;", text)
    if not match:
        raise ValueError(f"Missing numeric assignment: {name}")
    return np.array([[float(x) for x in row.split()] for row in match[1].strip().split(";") if row.strip()])


def read_map():
    folder = RAW / "epa_mg2"
    folder.mkdir(exist_ok=True)
    with ZipFile(RAW / "epa_2010_prius_mg2_emot_2023-04-06.zip") as z:
        assert z.testzip() is None
        inventory = [dict(file=i.filename, bytes=i.file_size, sha256=hashlib.sha256(z.read(i)).hexdigest())
                     for i in z.infolist()]
        for i in z.infolist():
            # Never extract unchecked paths or execute downloaded code.
            if Path(i.filename).name != i.filename:
                raise ValueError("Unexpected nested ZIP path")
            z.extract(i, folder)
    write_json(OUT / "epa_zip_inventory.json", inventory)
    text = next(folder.glob("7-*.m")).read_text(encoding="utf-8-sig")
    prefix = "mg.electric_power_W"
    speed = array_assignment(text, prefix + ".axis_1.breakpoints").ravel()
    torque = array_assignment(text, prefix + ".axis_2.breakpoints").ravel()
    power = array_assignment(text, prefix + ".table")
    limit = array_assignment(text, "mg.positive_torque_limit_Nm.table").ravel()
    negative = array_assignment(text, "mg.negative_torque_limit_Nm.table").ravel()
    rpm = np.round(speed * 60 / (2 * np.pi), 9)
    assert power.shape == (len(speed), len(torque)) == (14, 39)
    assert np.all(np.diff(speed) > 0) and np.all(np.diff(torque) > 0)
    assert np.isfinite(power).all() and np.allclose(negative, -limit)
    mechanical = speed[:, None] * torque[None, :]
    loss = power - mechanical
    assert np.min(loss) >= 0
    # Enforce both torque curve AND declared shaft-power ceiling.
    admissible = ((np.abs(torque)[None, :] <= limit[:, None] + 1e-9)
                  & (np.abs(mechanical) <= 60000 + 1e-8))
    with (OUT / "epa_emot_grid.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["speed_rpm", "torque_Nm", "electrical_power_W", "mechanical_power_W", "combined_loss_W", "inside_limits"])
        for i, n in enumerate(rpm):
            for j, t in enumerate(torque):
                w.writerow([n, t, power[i, j], mechanical[i, j], loss[i, j], bool(admissible[i, j])])
    # The source XLSX has sparse, separate motor/inverter DRIVING measurements.
    wb = load_workbook(next(folder.glob("3c*.xlsx")), data_only=True)
    motor, inv = wb["650V Motor Only"], wb["650V Inv Only"]
    paired = []
    counts = {}
    for sheet in wb:
        vals = [sheet.cell(r, c).value for r in range(8, 28) for c in range(3, 17)]
        counts[sheet.title] = dict(positive_torque_numeric_cells=sum(isinstance(x, (int, float)) for x in vals),
                                  positive_torque_missing_cells=sum(x is None for x in vals))
    for r in range(8, 28):
        for c in range(3, 17):
            em, ei = motor.cell(r, c).value, inv.cell(r, c).value
            if isinstance(em, (int, float)) and isinstance(ei, (int, float)):
                n, t = motor.cell(6, c).value, motor.cell(r, 2).value
                assert 0 < em <= 100 and 0 < ei <= 100
                paired.append(dict(speed_rpm=n, torque_Nm=t, motor_efficiency_norm=em/100,
                                   inverter_efficiency_norm=ei/100, combined_efficiency_norm=em*ei/10000))
    with (OUT / "ornl_paired_drive_efficiency.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(paired[0])); w.writeheader(); w.writerows(paired)
    interp = RegularGridInterpolator((rpm, torque), power, bounds_error=True)
    assert np.isclose(interp([[1000, 20]])[0], power[1, np.where(torque == 20)[0][0]])
    try:
        interp([[13500, 20]])
        raise AssertionError("Extrapolation must fail")
    except ValueError:
        pass
    # Sample XLSX is an incomplete demonstration, not the full grid.
    sample = load_workbook(next(folder.glob("6-*.xlsx")), data_only=True).active
    rows = [r for r in sample.iter_rows(min_row=6, values_only=True) if isinstance(r[0], (int, float))]
    for row in rows:
        assert np.isclose(interp([[row[0], row[1]]])[0]/1000, row[3], atol=1e-8)
    audit = dict(grid_shape=list(power.shape), speed_rpm_min=float(rpm[0]), speed_rpm_max=float(rpm[-1]),
                 torque_Nm_min=float(torque[0]), torque_Nm_max=float(torque[-1]),
                 missing_grid_cells=int(np.isnan(power).sum()), admissible_grid_cells=int(admissible.sum()),
                 min_combined_loss_W=float(loss.min()), raw_xlsx=counts, paired_drive_points=len(paired),
                 sample_xlsx_rows=len(rows), sample_speeds_rpm=sorted(set(r[0] for r in rows)),
                 regen_status="EPA estimated: motoring loss mirrored with factor 1.05; not measured regen",
                 speed_gap="Declared max 13500 rpm; tabulated map ends at 13000 rpm; no extrapolation",
                 voltage_V=650, reference_gear_ratio_not_to_apply=2.636,
                 independent_motor_winding_loss_available=False,
                 independent_motor_drive_efficiency_available=True,
                 revs_constructor_execution="NOT EXECUTED: ALPHA/REVS classes not bundled")
    write_json(OUT / "epa_map_audit.json", audit)
    mapdata = dict(speed_rpm=rpm.tolist(), torque_Nm=torque.tolist(), electrical_power_W=power.tolist(),
                   positive_torque_limit_Nm=limit.tolist(), negative_torque_limit_Nm=negative.tolist())
    write_json(OUT / "epa_map_numeric.json", mapdata)
    savemat(OUT / "epa_map_numeric.mat", {k:np.array(v) for k,v in mapdata.items()})
    return audit


def flatten(obj, prefix=""):
    for key, value in obj.items():
        name = f"{prefix}.{key}" if prefix else key
        if isinstance(value, dict):
            yield from flatten(value, name)
        else:
            yield name, value


def validate_config(cfg):
    """Reject obsolete topology and physically invalid M0 screening inputs."""
    assert cfg["metadata"]["version"] == "v0.4"
    v, d, b = cfg["vehicle"], cfg["drive"], cfg["brakes"]
    assert d["driven_axle"] == "rear" and d["independent_motor_count"] == 2
    assert d["differential_present"] is False
    assert b["rear_rotor_count"] == 2 and b["front_conventional_service_brakes_retained"] is True
    assert b["rotor_speed_to_own_wheel_speed_ratio"] == 1
    assert all(v[k] > 0 for k in ["mass_kg", "wheelbase_m", "wheel_effective_rolling_radius_m", "gravity_m_s2"])
    assert 0 < v["static_rear_axle_fraction"] < 1 and 0 < v["cg_height_m"] < v["wheelbase_m"]
    assert d["aggregate_motor_peak_power_kw"] == 2*d["rear_motor_per_side"]["peak_mechanical_power_kw"]
    for key in ["peak_drive_efficiency_screening", "peak_backdrive_efficiency_screening"]:
        assert 0 < d["rear_gearbox_per_side"][key] <= 1
    assert d["rear_gearbox_per_side"]["reduction_motor_to_wheel"] > 0
    assert d["rear_motor_per_side"]["configured_traction_torque_limit_nm"] <= d["rear_motor_per_side"]["physical_peak_torque_nm"]
    for z in cfg["cases"]["stop_deceleration_g_scan"]:
        assert z > 0 and 0 < v["static_rear_axle_fraction"]-z*v["cg_height_m"]/v["wheelbase_m"] < 1
    assert all(x >= 0 for x in cfg["cases"]["total_battery_regen_acceptance_kw_screening"])


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    cfg = yaml.safe_load((ROOT / "configs/baseline_vehicle.yaml").read_text(encoding="utf-8"))
    validate_config(cfg)
    assert cfg["metadata"]["version"] == "v0.4"
    assert cfg["drive"]["driven_axle"] == "rear"
    assert cfg["drive"]["independent_motor_count"] == 2
    assert cfg["drive"]["differential_present"] is False
    assert cfg["brakes"]["rotor_speed_to_own_wheel_speed_ratio"] == 1
    v, d = cfg["vehicle"], cfg["drive"]
    assert v["mass_kg"] > 0 and v["wheelbase_m"] > 0 and v["wheel_effective_rolling_radius_m"] > 0
    assert 0 < v["static_rear_axle_fraction"] < 1 and 0 < v["cg_height_m"] < v["wheelbase_m"]
    assert d["aggregate_motor_peak_power_kw"] == 2*d["rear_motor_per_side"]["peak_mechanical_power_kw"]
    for z in cfg["cases"]["stop_deceleration_g_scan"]:
        assert 0 < v["static_rear_axle_fraction"]-z*v["cg_height_m"]/v["wheelbase_m"] < 1
    write_json(OUT / "baseline_vehicle.json", cfg)
    evidence = {
        "vehicle.wheelbase_m": ("S1", "VW_MY21_reference_only", "VW MY21 PDF p19 dimension L101; visually verified"),
        "vehicle.drag_coefficient": ("S3", "UNVERIFIED_origin_404_search_index_only", "VW technical-data-6249; retain screening assumption"),
        "vehicle.tire_reference": ("S1", "OEM_reference_only", "VW MY23 PDF p5"),
        "vehicle.maximum_vehicle_speed_kmh_target": ("S3", "design_target_based_on_OEM", "VW MY23 PDF p9"),
        "vehicle.tire_nominal_unloaded_radius_m": ("S3", "derived_geometry", "215/55R18 nominal dimensions"),
    }
    rows = []
    for name, value in flatten(cfg):
        grade, status, source = evidence.get(name, ("S3", "concept_or_unverified", "V0.4 configuration; see report"))
        if "rear_motor_per_side" in name and name.rsplit(".",1)[-1] in ["peak_mechanical_power_kw", "physical_peak_torque_nm", "max_mechanical_speed_rpm", "stator_outer_diameter_m"]:
            grade, status, source = "S1", "reference_motor_only", "Toyota launch; ORNL/TM-2010/253 Tables 2.1,2.7; Sec3.3"
        if name.startswith("brakes.rear_rotor_per_side.") and name.rsplit(".",1)[-1] in ["outside_diameter_m","thickness_m","disc_overall_height_m","nominal_center_bore_m"]:
            grade, status, source = "S1", "catalogue_geometry_only", "Brembo 09.9908.21"
        if "map_file" in name and value is not None:
            grade, status, source = "S1", "data_downloaded_not_validated_for_concept", "EPA package 2023-04"
        if value is None:
            grade, status, source = "unknown", "MISSING", "not supplied"
        if name.startswith(("metadata.","provenance.")) or isinstance(value, (str,bool)) and "map_file" not in name:
            if name not in evidence:
                grade, status = "N/A", "metadata_or_design_declaration"
        unit = next((u for suffix,u in [("_m2K_W","m2 K/W"),("_kg_s","kg/s"),("_m_s2","m/s2"),
                                       ("_m_bounds","m"),("_deg_bounds","deg"),("_rpm","rpm"),
                                       ("_kmh_target","km/h"),("_kmh","km/h"),("_m2","m2"),
                                       ("_kg","kg"),("_kw","kW"),("_nm","N m"),("_m","m"),
                                       ("_c","degC"),("_N","N"),("_s","s"),("_g","g units")]
                     if name.endswith(suffix)), "1" if isinstance(value,(float,int,list)) else "N/A_or_unknown")
        rows.append(dict(id=name, priority="P0" if value is None else "P1", value=json.dumps(value,ensure_ascii=False), unit=unit,
                         status=status, source_grade=grade, source=source, uncertainty="see M0 report; no measured concept data",
                         blocks_stage="M1_or_later" if value is None else "review", action="obtain evidence" if value is None else "reference or screening only"))
    with (OUT / "parameter_audit.csv").open("w", newline="", encoding="utf-8-sig") as f:
        w=csv.DictWriter(f, fieldnames=list(rows[0])); w.writeheader(); w.writerows(rows)
    audit = read_map()
    meta=dict(base_commit=subprocess.check_output(["git","rev-parse","HEAD"],cwd=ROOT,text=True).strip(),
              config_sha256=hashlib.sha256((ROOT/"configs/baseline_vehicle.yaml").read_text(encoding="utf-8").encode("utf-8")).hexdigest(),
              config_hash_encoding="UTF-8 with LF line endings; raw-source hashes remain byte-exact",
              python=platform.python_version(), platform=platform.platform(), null_parameters=sum(r["status"]=="MISSING" for r in rows),
              scope="M0 only; M1-M7 not authorized", map_audit=audit)
    write_json(OUT / "audit_metadata.json", meta)
    print(json.dumps(meta,ensure_ascii=False,indent=2))


if __name__ == "__main__":
    main()
