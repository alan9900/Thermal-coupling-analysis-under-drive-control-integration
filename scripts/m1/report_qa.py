"""Pandoc style template and Word-export visual QA for the M1 report."""
import json
import sys
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZipFile

import pymupdf
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT=Path(__file__).resolve().parents[2]


def reference():
    d=Document(ROOT/'reports/M0_reference.docx')
    for sec in d.sections:
        for p in sec.footer.paragraphs:
            for run in p.runs:
                if 'M0' in run.text:
                    run.text=run.text.replace('M0 参数审查','M1 多热源模型')
    d.save(ROOT/'reports/M1_reference.docx')


def render():
    import win32com.client
    path=ROOT/'reports/M1_report.docx'
    d=Document(path)
    for table in d.tables:
        keep_short=table.cell(0,0).text=='减速度'
        for ri,row in enumerate(table.rows):
            trpr=row._tr.get_or_add_trPr()
            if trpr.find(qn('w:cantSplit')) is None:
                trpr.append(OxmlElement('w:cantSplit'))
            for cell in row.cells:
                if ri==0:
                    shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'EAF0F7'); cell._tc.get_or_add_tcPr().append(shade)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.05
                    if keep_short and ri<len(table.rows)-1:
                        p.paragraph_format.keep_with_next=True
                    for run in p.runs:
                        run.font.size=Pt(9.5)
                        if ri==0: run.font.bold=True
    d.core_properties.language='zh-CN'; d.save(path)
    with ZipFile(path) as z:
        tree=ET.fromstring(z.read('word/document.xml'))
        equations=len(tree.findall('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath'))
        assert equations>=20
        assert len([n for n in z.namelist() if n.startswith('word/media/')])==2
    word=win32com.client.DispatchEx('Word.Application')
    word.Visible=False; word.DisplayAlerts=0; word.AutomationSecurity=3; doc=None
    try:
        doc=word.Documents.Open(str(path),ReadOnly=True,AddToRecentFiles=False)
        doc.ExportAsFixedFormat(str(ROOT/'reports/M1_report.pdf'),17)
    finally:
        if doc is not None: doc.Close(False)
        word.Quit()
    folder=ROOT/'tmp/m1_qa/report'; folder.mkdir(parents=True,exist_ok=True)
    pdf=pymupdf.open(ROOT/'reports/M1_report.pdf')
    text='\n'.join(p.get_text() for p in pdf)
    assert 'PARTIAL' in text and '75.87' in text and 'M0-v1.0' in text
    for i,page in enumerate(pdf):
        page.get_pixmap(matrix=pymupdf.Matrix(1.25,1.25)).save(folder/f'page_{i+1:02d}.png')
    result=dict(word_export='PASS',editable_math_objects=equations,pdf_pages=len(pdf),embedded_figures=2,
                post_pandoc_changes='formatting only: table headers, row splitting, spacing and language metadata',
                visual_inspection='PENDING')
    (ROOT/'results/m1/document_checks.json').write_text(json.dumps(result,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(result,indent=2))


if __name__=='__main__':
    {'reference':reference,'render':render}[sys.argv[1]]()
