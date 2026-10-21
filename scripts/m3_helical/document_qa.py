"""Pandoc OMML + actual Word PDF render, then explicit visual QA record."""
import sys,json,hashlib
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET
import pymupdf
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'results/m3_helical';BASE='M3_helical_revision'
def main():
    path=ROOT/'reports'/f'{BASE}.docx'
    if '--record-visual' in sys.argv:
        p=OUT/'document_checks.json';r=json.loads(p.read_text())
        for ext,h in r['sha256'].items():assert hashlib.sha256((ROOT/'reports'/f'{BASE}.{ext}').read_bytes()).hexdigest()==h
        r.update(visual_inspection='PASS',pages_inspected=list(range(1,r['pdf_pages']+1)))
        p.write_text(json.dumps(r,indent=2)+'\n');return
    d=Document(path)
    for table in d.tables:
        for i,row in enumerate(table.rows):
            props=row._tr.get_or_add_trPr();props.append(OxmlElement('w:cantSplit'))
            if i==0:props.append(OxmlElement('w:tblHeader'))
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(3)
                    p.paragraph_format.keep_with_next=i<len(table.rows)-1
                    for run in p.runs:run.font.size=Pt(9)
    d.core_properties.language='zh-CN';d.save(path)
    with ZipFile(path) as z:
        xml=ET.fromstring(z.read('word/document.xml'))
        math=len(xml.findall('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath'))
        images=len([n for n in z.namelist() if n.startswith('word/media/')])
        assert math==5 and images==5,(math,images)
    import win32com.client
    app=win32com.client.DispatchEx('Word.Application');app.Visible=False;app.DisplayAlerts=0;app.AutomationSecurity=3;doc=None
    try:
        doc=app.Documents.Open(str(path),ReadOnly=True,AddToRecentFiles=False)
        doc.ExportAsFixedFormat(str(ROOT/'reports'/f'{BASE}.pdf'),17)
    finally:
        if doc is not None:doc.Close(False)
        app.Quit()
    pdf=pymupdf.open(ROOT/'reports'/f'{BASE}.pdf');folder=ROOT/'tmp/m3_helical_qa';folder.mkdir(exist_ok=True)
    for i,p in enumerate(pdf):p.get_pixmap(matrix=pymupdf.Matrix(1.1,1.1)).save(folder/f'page_{i+1:02d}.png')
    text='\n'.join(p.get_text() for p in pdf)
    for word in ['BLOCKED','PARTIAL','NOT EXECUTED','6624.258']:assert word in text,word
    r=dict(status='PASS_RENDER_STRUCTURE',editable_math_objects=math,figures=images,pdf_pages=len(pdf),visual_inspection='PENDING',
        sha256={ext:hashlib.sha256((ROOT/'reports'/f'{BASE}.{ext}').read_bytes()).hexdigest() for ext in ['md','docx','pdf']})
    (OUT/'document_checks.json').write_text(json.dumps(r,indent=2)+'\n');print(r)
if __name__=='__main__':main()
